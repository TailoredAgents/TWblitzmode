"""
Prospect Ingestion Service

Handles CSV/XLSX file uploads and validation for the corporate warm-intro workflow.
Processes prospect data, validates required fields, and queues executive lookup jobs.

Features:
- Support for CSV and XLSX file formats
- Comprehensive data validation and sanitization
- Duplicate detection with organization isolation
- LinkedIn URL canonicalization
- Batch processing with progress tracking
- Error reporting and data quality metrics
"""

import os
import re
import csv
import io
import json
import logging
import asyncio
from datetime import datetime, timezone
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import Dict, Any, List, Optional, Tuple, Union
import mimetypes
from dataclasses import dataclass, asdict
from enum import Enum

try:
    import pandas as pd  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    pd = None  # type: ignore

try:
    import validators  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    validators = None  # type: ignore
try:
    from fastapi import UploadFile, HTTPException
except ImportError:  # pragma: no cover - fallback for environments without FastAPI
    class HTTPException(RuntimeError):
        def __init__(self, status_code: int, detail: str):
            self.status_code = status_code
            self.detail = detail
            super().__init__(detail)

    class UploadFile:  # Minimal stub used only for type hints in optional environments
        def __init__(self, filename: str = "", file=None):
            self.filename = filename
            self._file = file

        async def read(self) -> bytes:
            if self._file is None:
                raise RuntimeError("UploadFile stub cannot read data without an underlying file object.")
            return self._file.read()

        async def seek(self, offset: int) -> None:
            if self._file is None:
                return
            self._file.seek(offset)

try:
    from docx import Document  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    Document = None

# Import Redis for real-time progress tracking
try:
    import redis
    REDIS_AVAILABLE = True
except ImportError:
    REDIS_AVAILABLE = False

logger = logging.getLogger(__name__)

class IngestionStatus(Enum):
    """Ingestion job status enumeration"""
    PENDING = "pending"
    PROCESSING = "processing"
    COMPLETED = "completed"
    FAILED = "failed"
    PARTIALLY_COMPLETED = "partially_completed"

@dataclass
class ProspectRecord:
    """Validated prospect record"""
    company: str
    full_name: str
    role: str = "executive"
    email: Optional[str] = None
    linkedin_url: Optional[str] = None
    location: Optional[str] = None
    tags: List[str] = None
    notes: Optional[str] = None
    source_row: int = 0
    id: Optional[int] = None

    def __post_init__(self):
        if self.tags is None:
            self.tags = []

@dataclass
class ValidationError:
    """Validation error record"""
    row: int
    field: str
    value: str
    error: str
    severity: str = "error"  # "error", "warning"

@dataclass
class IngestionResult:
    """Result of the ingestion process"""
    status: IngestionStatus
    total_rows: int
    valid_prospects: int
    invalid_prospects: int
    duplicates_found: int
    prospects_created: int
    lookup_jobs_queued: int
    validation_errors: List[ValidationError]
    processing_time_seconds: float
    cost_estimate_usd: float = 0.0
    prospects: List[ProspectRecord] = None

    def __post_init__(self):
        if self.prospects is None:
            self.prospects = []

class ProspectIngestionService:
    """Production-ready prospect ingestion service"""

    # Required CSV columns
    REQUIRED_COLUMNS = {"company", "executive_name"}

    # Optional columns with aliases
    COLUMN_MAPPING = {
        "executive_name": ["full_name", "name", "executive_name", "prospect_name"],
        "executive_first_name": ["first_name", "fname"],
        "executive_last_name": ["last_name", "lname", "surname"],
        "company": ["company", "company_name", "organization", "employer"],
        "role": ["role", "title", "position", "job_title"],
        "email": ["email", "email_address", "contact_email"],
        "linkedin_url": ["linkedin_url", "linkedin", "linkedin_profile", "profile_url"],
        "location": ["location", "city", "country", "region"],
        "notes": ["notes", "comments", "description"]
    }

    # Data validation patterns
    EMAIL_PATTERN = re.compile(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$')
    LINKEDIN_PATTERNS = [
        re.compile(r'linkedin\.com/in/([a-zA-Z0-9\-]+)/?', re.IGNORECASE),
        re.compile(r'linkedin\.com/pub/([a-zA-Z0-9\-]+)/?', re.IGNORECASE)
    ]

    def __init__(self):
        self.max_file_size_mb = int(os.getenv('MAX_UPLOAD_SIZE_MB', '50'))
        self.max_rows_per_batch = int(os.getenv('MAX_PROSPECTS_PER_BATCH', '1000'))

        # Redis connection for progress tracking
        self.redis_client = None
        if REDIS_AVAILABLE:
            try:
                redis_url = os.getenv('REDIS_URL', 'redis://localhost:6379')
                self.redis_client = redis.from_url(redis_url)
                logger.info("Redis connection established for progress tracking")
            except Exception as e:
                logger.warning(f"Failed to connect to Redis: {e}")

        # Progress tracking configuration
        self.enable_progress_tracking = os.getenv('ENABLE_PROGRESS_TRACKING', 'true').lower() == 'true'

    async def process_file_upload(
        self,
        file: UploadFile,
        organization_id: int,
        uploaded_by: int,
        tags: List[str] = None,
        job_id: str = None
    ) -> IngestionResult:
        """Process uploaded CSV/XLSX file and create prospects"""

        start_time = datetime.now(timezone.utc)

        # Initialize progress tracking
        if not job_id:
            job_id = f"ingest_{organization_id}_{int(start_time.timestamp())}"

        try:
            await self._update_progress(job_id, "validating_file", 0, "Validating file format and size")

            # Validate file
            await self._validate_file(file)

            await self._update_progress(job_id, "parsing_file", 10, "Parsing file content")

            # Parse file content
            raw_data = await self._parse_file(file)

            await self._update_progress(job_id, "validating_data", 25, f"Validating {len(raw_data)} prospect records")

            # Validate and clean data
            prospects, validation_errors = await self._validate_prospects(raw_data, tags or [], job_id)

            await self._update_progress(job_id, "checking_duplicates", 60, "Checking for existing prospects")

            # Check for duplicates within organization
            prospects, duplicates = await self._check_duplicates(prospects, organization_id)

            await self._update_progress(job_id, "creating_prospects", 75, f"Creating {len(prospects)} prospect records")

            # Create prospects in database
            created_prospects = await self._create_prospects(prospects, organization_id, uploaded_by)

            await self._update_progress(job_id, "queuing_jobs", 90, "Queuing executive lookup jobs")

            # Queue executive lookup jobs
            lookup_jobs = await self._queue_lookup_jobs(created_prospects, organization_id)

            # Calculate processing time
            processing_time = (datetime.now(timezone.utc) - start_time).total_seconds()

            # Determine final status
            status = self._determine_status(len(raw_data), len(created_prospects), validation_errors)

            await self._update_progress(job_id, "completed", 100, "Ingestion completed successfully")

            return IngestionResult(
                status=status,
                total_rows=len(raw_data),
                valid_prospects=len(prospects),
                invalid_prospects=len(validation_errors),
                duplicates_found=len(duplicates),
                prospects_created=len(created_prospects),
                lookup_jobs_queued=len(lookup_jobs),
                validation_errors=validation_errors,
                processing_time_seconds=processing_time,
                cost_estimate_usd=len(lookup_jobs) * 0.02,  # Estimated cost per lookup
                prospects=prospects
            )

        except Exception as e:
            logger.error(f"Ingestion failed for organization {organization_id}: {e}")
            await self._update_progress(job_id, "failed", 0, f"Ingestion failed: {str(e)}")
            return IngestionResult(
                status=IngestionStatus.FAILED,
                total_rows=0,
                valid_prospects=0,
                invalid_prospects=0,
                duplicates_found=0,
                prospects_created=0,
                lookup_jobs_queued=0,
                validation_errors=[ValidationError(0, "file", str(file.filename), str(e))],
                processing_time_seconds=(datetime.now(timezone.utc) - start_time).total_seconds(),
                prospects=[]
            )

    async def process_file(
        self,
        file_path: Union[str, Path],
        organization_id: int,
        uploaded_by: int,
        tags: Optional[List[str]] = None,
        job_id: Optional[str] = None
    ) -> IngestionResult:
        """
        Process a file that already exists on disk.

        This is a thin wrapper that hydrates a Starlette UploadFile so we can
        reuse the validated upload pipeline used by the API layer.
        """

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Ingestion file not found: {file_path}")

        content = await asyncio.to_thread(path.read_bytes)
        spooled: SpooledTemporaryFile = SpooledTemporaryFile(max_size=len(content) + 1024)
        spooled.write(content)
        spooled.seek(0)

        upload = UploadFile(
            filename=path.name,
            file=spooled,
            content_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        )

        try:
            result = await self.process_file_upload(
                file=upload,
                organization_id=organization_id,
                uploaded_by=uploaded_by,
                tags=tags or [],
                job_id=job_id
            )
        finally:
            await upload.close()

        return result

    async def _validate_file(self, file: UploadFile) -> None:
        """Validate file size and format"""

        # Check file size
        file_size_mb = len(await file.read()) / (1024 * 1024)
        await file.seek(0)  # Reset file pointer

        if file_size_mb > self.max_file_size_mb:
            raise HTTPException(
                status_code=413,
                detail=f"File size ({file_size_mb:.1f}MB) exceeds limit ({self.max_file_size_mb}MB)"
            )

        # Check file format
        filename = file.filename.lower()
        if not filename.endswith(('.csv', '.xlsx', '.xls', '.docx')):
            raise HTTPException(
                status_code=400,
                detail="Unsupported file format. Please upload CSV or Excel files."
            )

    async def _parse_file(self, file: UploadFile) -> List[Dict[str, Any]]:
        """Parse CSV or Excel file into list of dictionaries"""

        try:
            content = await file.read()
            filename = file.filename.lower()

            if filename.endswith('.csv'):
                # Parse CSV
                text_content = content.decode('utf-8-sig')  # Handle BOM
                csv_reader = csv.DictReader(io.StringIO(text_content))
                data = [row for row in csv_reader]

            elif filename.endswith(('.xlsx', '.xls')):
                if pd is None:
                    raise HTTPException(
                        status_code=400,
                        detail="Excel ingestion requires the pandas package. Please install pandas to process Excel files."
                    )
                # Parse Excel
                df = pd.read_excel(io.BytesIO(content))
                data = df.to_dict('records')

            elif filename.endswith('.docx'):
                data = self._parse_docx(io.BytesIO(content))

            else:
                raise ValueError("Unsupported file format")

            if len(data) > self.max_rows_per_batch:
                raise HTTPException(
                    status_code=400,
                    detail=f"File contains {len(data)} rows, maximum allowed is {self.max_rows_per_batch}"
                )

            logger.info(f"Parsed {len(data)} rows from {file.filename}")
            return data

        except UnicodeDecodeError:
            raise HTTPException(status_code=400, detail="File encoding not supported. Please use UTF-8.")
        except Exception as e:
            if pd is not None and isinstance(e, pd.errors.EmptyDataError):
                raise HTTPException(status_code=400, detail="File is empty or contains no data.")
            raise HTTPException(status_code=400, detail=f"Failed to parse file: {str(e)}")

    def _parse_docx(self, stream: io.BytesIO) -> List[Dict[str, Any]]:
        """Extract company rows from a DOCX document."""

        if Document is None:
            raise HTTPException(
                status_code=400,
                detail="DOCX ingestion requires the python-docx package. Please contact support."
            )

        try:
            document = Document(stream)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Unable to read DOCX file: {exc}")

        raw_entries: List[str] = []

        # Extract text from tables first (structured content)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    entry = cell.text.strip()
                    if entry:
                        raw_entries.append(entry)

        # Extract text from paragraphs
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                raw_entries.append(text)

        companies = self._normalize_company_entries(raw_entries)

        if not companies:
            raise HTTPException(
                status_code=400,
                detail="No company names detected in the document. Provide one company name per line or table row."
            )

        return [
            {
                "company": company,
                "executive_name": "TBD",
                "notes": "Imported from DOCX company list"
            }
            for company in companies
        ]

    def _normalize_company_entries(self, entries: List[str]) -> List[str]:
        """Normalize free-form company input into a clean, unique list."""

        normalized: List[str] = []
        seen = set()

        separators = [",", ";", "|", "\n"]

        for entry in entries:
            # Replace bullet characters and trim whitespace
            cleaned = entry.replace("•", " ").replace("–", " ").replace("—", " ").strip()
            if not cleaned:
                continue

            # Split on separators while preserving meaningful text
            segments = [cleaned]
            for sep in separators:
                new_segments = []
                for segment in segments:
                    new_segments.extend(part.strip() for part in segment.split(sep))
                segments = new_segments

            for segment in segments:
                if not segment:
                    continue

                # Remove leading numbering/bullets (e.g., "1. Company")
                segment = re.sub(r"^[\d\.\-\)\s]+", "", segment).strip()

                if segment and segment.lower() not in seen:
                    seen.add(segment.lower())
                    normalized.append(segment)

        return normalized

    async def _validate_prospects(
        self,
        raw_data: List[Dict[str, Any]],
        default_tags: List[str],
        job_id: str = None
    ) -> Tuple[List[ProspectRecord], List[ValidationError]]:
        """Validate and clean prospect data"""

        prospects = []
        errors = []

        # Check for required columns
        if not raw_data:
            errors.append(ValidationError(0, "file", "", "File contains no data rows"))
            return prospects, errors

        # Normalize column names
        first_row = raw_data[0]
        column_mapping = self._create_column_mapping(list(first_row.keys()))

        # Check if required columns are present
        missing_columns = []
        for required in self.REQUIRED_COLUMNS:
            if required not in column_mapping:
                missing_columns.append(required)

        if missing_columns:
            errors.append(ValidationError(
                0, "headers", str(missing_columns),
                f"Missing required columns: {', '.join(missing_columns)}"
            ))
            return prospects, errors

        # Validate each row
        total_rows = len(raw_data)
        for row_idx, row_data in enumerate(raw_data, 1):
            try:
                prospect = await self._validate_single_prospect(
                    row_data, row_idx, column_mapping, default_tags
                )
                if prospect:
                    prospects.append(prospect)

                # Update progress every 50 rows
                if job_id and row_idx % 50 == 0:
                    progress = 25 + int((row_idx / total_rows) * 35)  # 25-60% range
                    await self._update_progress(
                        job_id, "validating_data", progress,
                        f"Validated {row_idx}/{total_rows} prospects"
                    )

            except ValidationError as ve:
                errors.append(ve)
            except Exception as e:
                errors.append(ValidationError(
                    row_idx, "row", str(row_data), f"Unexpected validation error: {str(e)}"
                ))

        logger.info(f"Validated {len(prospects)} prospects with {len(errors)} errors")
        return prospects, errors

    def _create_column_mapping(self, columns: List[str]) -> Dict[str, str]:
        """Create mapping from standard names to actual column names"""

        mapping = {}
        columns_lower = [col.lower().strip() for col in columns]

        for standard_name, aliases in self.COLUMN_MAPPING.items():
            for alias in aliases:
                if alias.lower() in columns_lower:
                    original_col = columns[columns_lower.index(alias.lower())]
                    mapping[standard_name] = original_col
                    break

        return mapping

    async def _validate_single_prospect(
        self,
        row_data: Dict[str, Any],
        row_idx: int,
        column_mapping: Dict[str, str],
        default_tags: List[str]
    ) -> Optional[ProspectRecord]:
        """Validate a single prospect record"""

        # Extract mapped values
        company = self._get_mapped_value(row_data, column_mapping, "company", "").strip()

        # Handle name fields - try combined name first, then first+last
        full_name = self._get_mapped_value(row_data, column_mapping, "executive_name", "").strip()
        if not full_name:
            first_name = self._get_mapped_value(row_data, column_mapping, "executive_first_name", "").strip()
            last_name = self._get_mapped_value(row_data, column_mapping, "executive_last_name", "").strip()
            full_name = f"{first_name} {last_name}".strip()

        # Required field validation
        if not company:
            raise ValidationError(row_idx, "company", company, "Company name is required")

        if not full_name:
            raise ValidationError(row_idx, "executive_name", full_name, "Executive name is required")

        if len(company) > 255:
            raise ValidationError(row_idx, "company", company, "Company name too long (max 255 characters)")

        if len(full_name) > 255:
            raise ValidationError(row_idx, "executive_name", full_name, "Executive name too long (max 255 characters)")

        # Optional field validation
        role = self._get_mapped_value(row_data, column_mapping, "role", "executive").strip()
        email = self._get_mapped_value(row_data, column_mapping, "email", "").strip()
        linkedin_url = self._get_mapped_value(row_data, column_mapping, "linkedin_url", "").strip()
        location = self._get_mapped_value(row_data, column_mapping, "location", "").strip()
        notes = self._get_mapped_value(row_data, column_mapping, "notes", "").strip()

        # Email validation
        if email and not self.EMAIL_PATTERN.match(email):
            raise ValidationError(row_idx, "email", email, "Invalid email format")

        # LinkedIn URL validation and canonicalization
        linkedin_canonical = None
        if linkedin_url:
            linkedin_canonical = self._canonicalize_linkedin_url(linkedin_url)
            if not linkedin_canonical:
                raise ValidationError(row_idx, "linkedin_url", linkedin_url, "Invalid LinkedIn URL format")

        # Create prospect record
        return ProspectRecord(
            company=company,
            full_name=full_name,
            role=role,
            email=email if email else None,
            linkedin_url=linkedin_canonical,
            location=location if location else None,
            tags=default_tags.copy(),
            notes=notes if notes else None,
            source_row=row_idx
        )

    def _get_mapped_value(
        self,
        row_data: Dict[str, Any],
        column_mapping: Dict[str, str],
        field_name: str,
        default: Any = None
    ) -> Any:
        """Get value from row using column mapping"""

        if field_name in column_mapping:
            actual_column = column_mapping[field_name]
            value = row_data.get(actual_column, default)
            # Handle NaN values from pandas
            if pd is not None and pd.isna(value):
                return default
            return str(value) if value is not None else default

        return default

    def _canonicalize_linkedin_url(self, url: str) -> Optional[str]:
        """Canonicalize LinkedIn URL to standard format"""

        url = url.strip()

        # Add https if missing
        if not url.startswith(('http://', 'https://')):
            url = 'https://' + url

        # Extract username from various LinkedIn URL formats
        for pattern in self.LINKEDIN_PATTERNS:
            match = pattern.search(url)
            if match:
                username = match.group(1)
                return f"https://www.linkedin.com/in/{username}/"

        # Check if it's a valid LinkedIn URL but not in expected format
        if 'linkedin.com' in url.lower():
            if validators is None:
                return url
            if validators.url(url):
                return url  # Return as-is if it's a valid LinkedIn URL

        return None

    async def _check_duplicates(
        self,
        prospects: List[ProspectRecord],
        organization_id: int
    ) -> Tuple[List[ProspectRecord], List[ProspectRecord]]:
        """Check for duplicates within organization and remove them"""

        from ..api.db_core import get_conn, query

        unique_prospects = []
        duplicates = []

        with get_conn() as conn:
            for prospect in prospects:
                # Check if prospect already exists in organization
                existing = query(conn, """
                    SELECT id, full_name, company
                    FROM prospects
                    WHERE organization_id = ? AND company = ? AND full_name = ?
                """, (organization_id, prospect.company, prospect.full_name))

                if existing:
                    duplicates.append(prospect)
                    logger.debug(f"Duplicate found: {prospect.full_name} at {prospect.company}")
                else:
                    unique_prospects.append(prospect)

        logger.info(f"Filtered {len(duplicates)} duplicates, {len(unique_prospects)} unique prospects remain")
        return unique_prospects, duplicates

    async def _create_prospects(
        self,
        prospects: List[ProspectRecord],
        organization_id: int,
        created_by: int
    ) -> List[int]:
        """Create prospect records in database"""

        from ..api.db_core import get_conn, execute

        created_ids = []

        with get_conn() as conn:
            for prospect in prospects:
                try:
                    prospect_id = execute(conn, """
                        INSERT INTO prospects (
                            organization_id, company, full_name, role, email,
                            linkedin_url, location, tags, notes, status, source, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending_lookup', 'csv_import', CURRENT_TIMESTAMP)
                    """, (
                        organization_id, prospect.company, prospect.full_name,
                        prospect.role, prospect.email, prospect.linkedin_url,
                        prospect.location, json.dumps(prospect.tags), prospect.notes
                    ))

                    prospect.id = prospect_id
                    created_ids.append(prospect_id)

                except Exception as e:
                    logger.error(f"Failed to create prospect {prospect.full_name}: {e}")
                    # Continue with other prospects

        logger.info(f"Created {len(created_ids)} prospect records")
        return created_ids

    async def _queue_lookup_jobs(
        self,
        prospect_ids: List[int],
        organization_id: int
    ) -> List[str]:
        """Queue executive lookup jobs for created prospects"""

        from ..api.db_core import get_conn, execute, create_job_idempotent

        job_keys = []

        with get_conn() as conn:
            for prospect_id in prospect_ids:
                try:
                    # Create idempotent job
                    idempotency_key = f"executive_lookup:{organization_id}:{prospect_id}"

                    # Check if job already exists
                    from ..api.db_core import query
                    existing = query(conn, """
                        SELECT id FROM job_idempotency
                        WHERE organization_id = ? AND idempotency_key = ?
                    """, (organization_id, idempotency_key))

                    if not existing:
                        create_job_idempotent(
                            conn, organization_id, idempotency_key,
                            prospect_id, "executive_lookup"
                        )
                        job_keys.append(idempotency_key)

                except Exception as e:
                    logger.error(f"Failed to queue lookup job for prospect {prospect_id}: {e}")

        logger.info(f"Queued {len(job_keys)} executive lookup jobs")
        return job_keys

    def _determine_status(
        self,
        total_rows: int,
        created_prospects: int,
        validation_errors: List[ValidationError]
    ) -> IngestionStatus:
        """Determine final ingestion status"""

        if created_prospects == 0:
            return IngestionStatus.FAILED
        elif len(validation_errors) > 0:
            return IngestionStatus.PARTIALLY_COMPLETED
        else:
            return IngestionStatus.COMPLETED

    async def get_ingestion_summary(self, organization_id: int, days: int = 30) -> Dict[str, Any]:
        """Get ingestion summary for organization dashboard"""

        from ..api.db_core import get_conn, query

        with get_conn() as conn:
            # Get recent imports
            recent_imports = query(conn, """
                SELECT
                    COUNT(*) as total_prospects,
                    COUNT(CASE WHEN status = 'pending_lookup' THEN 1 END) as pending_lookup,
                    COUNT(CASE WHEN status = 'lookup_failed' THEN 1 END) as lookup_failed,
                    COUNT(CASE WHEN status = 'ready_for_mutuals' THEN 1 END) as ready_for_mutuals,
                    DATE(created_at) as import_date
                FROM prospects
                WHERE organization_id = ?
                AND source = 'csv_import'
                AND created_at >= DATE('now', '-{} days')
                GROUP BY DATE(created_at)
                ORDER BY import_date DESC
            """.format(days), (organization_id,))

            return {
                "organization_id": organization_id,
                "period_days": days,
                "daily_imports": [dict(row) for row in recent_imports],
                "total_imported": sum(row["total_prospects"] for row in recent_imports),
                "status_breakdown": {
                    "pending_lookup": sum(row["pending_lookup"] for row in recent_imports),
                    "lookup_failed": sum(row["lookup_failed"] for row in recent_imports),
                    "ready_for_mutuals": sum(row["ready_for_mutuals"] for row in recent_imports)
                }
            }

    async def _update_progress(
        self,
        job_id: str,
        stage: str,
        progress_percent: int,
        message: str
    ):
        """Update ingestion progress in Redis for real-time tracking"""

        if not self.enable_progress_tracking or not self.redis_client:
            return

        try:
            progress_data = {
                "job_id": job_id,
                "stage": stage,
                "progress_percent": progress_percent,
                "message": message,
                "timestamp": datetime.now(timezone.utc).isoformat()
            }

            # Store progress with 1 hour expiry
            self.redis_client.setex(
                f"ingestion_progress:{job_id}",
                3600,
                json.dumps(progress_data)
            )

        except Exception as e:
            logger.warning(f"Failed to update progress for job {job_id}: {e}")

    async def get_ingestion_progress(self, job_id: str) -> Optional[Dict[str, Any]]:
        """Get current ingestion progress"""

        if not self.redis_client:
            return None

        try:
            progress_data = self.redis_client.get(f"ingestion_progress:{job_id}")
            if progress_data:
                return json.loads(progress_data)

        except Exception as e:
            logger.warning(f"Failed to get progress for job {job_id}: {e}")

        return None

    async def validate_csv_headers(self, file: UploadFile) -> Dict[str, Any]:
        """Validate CSV headers without processing the full file"""

        try:
            # Read just the first few lines to check headers
            content = await file.read(1024)  # Read first 1KB
            await file.seek(0)  # Reset file pointer

            # Parse headers
            if file.filename.lower().endswith('.csv'):
                text_content = content.decode('utf-8-sig')
                lines = text_content.split('\n')
                if lines:
                    headers = lines[0].split(',')
                    headers = [h.strip().strip('"') for h in headers]
                else:
                    return {"valid": False, "error": "File appears to be empty"}
            else:
                return {"valid": False, "error": "Only CSV header validation supported"}

            # Check column mapping
            column_mapping = self._create_column_mapping(headers)

            # Check required columns
            missing_required = []
            for required in self.REQUIRED_COLUMNS:
                if required not in column_mapping:
                    missing_required.append(required)

            # Check recognized columns
            recognized_columns = list(column_mapping.keys())
            unrecognized_columns = [h for h in headers if h not in [column_mapping.get(k) for k in column_mapping.keys()]]

            return {
                "valid": len(missing_required) == 0,
                "headers": headers,
                "missing_required": missing_required,
                "recognized_columns": recognized_columns,
                "unrecognized_columns": unrecognized_columns,
                "column_mapping": column_mapping
            }

        except Exception as e:
            return {"valid": False, "error": f"Failed to validate headers: {str(e)}"}

    async def get_template_csv(self) -> str:
        """Generate a CSV template with proper headers"""

        headers = [
            "company",
            "executive_name",
            "role",
            "email",
            "linkedin_url",
            "location",
            "notes"
        ]

        # Create sample data
        sample_data = [
            {
                "company": "Acme Corp",
                "executive_name": "John Smith",
                "role": "CEO",
                "email": "john.smith@acme.com",
                "linkedin_url": "https://www.linkedin.com/in/johnsmith",
                "location": "San Francisco, CA",
                "notes": "Referred by Jane Doe"
            },
            {
                "company": "Tech Startup Inc",
                "executive_name": "Sarah Johnson",
                "role": "CTO",
                "email": "",
                "linkedin_url": "https://www.linkedin.com/in/sarah-johnson-tech",
                "location": "Austin, TX",
                "notes": ""
            }
        ]

        # Generate CSV
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(sample_data)

        return output.getvalue()

    async def reprocess_failed_prospects(
        self,
        organization_id: int,
        prospect_ids: List[int] = None
    ) -> Dict[str, Any]:
        """Reprocess prospects that failed validation or lookup"""

        from ..api.db_core import get_conn, query, execute

        try:
            with get_conn() as conn:
                # Get failed prospects
                if prospect_ids:
                    prospects = query(conn, """
                        SELECT id, company, full_name, role, status
                        FROM prospects
                        WHERE organization_id = ? AND id IN ({})
                        AND status IN ('lookup_failed', 'validation_failed')
                    """.format(','.join('?' * len(prospect_ids))),
                    [organization_id] + prospect_ids)
                else:
                    prospects = query(conn, """
                        SELECT id, company, full_name, role, status
                        FROM prospects
                        WHERE organization_id = ?
                        AND status IN ('lookup_failed', 'validation_failed')
                        LIMIT 100
                    """, (organization_id,))

                if not prospects:
                    return {"reprocessed": 0, "message": "No failed prospects found"}

                # Reset status to pending_lookup
                prospect_ids_to_reprocess = [p['id'] for p in prospects]
                execute(conn, """
                    UPDATE prospects
                    SET status = 'pending_lookup', processed_at = NULL
                    WHERE id IN ({})
                """.format(','.join('?' * len(prospect_ids_to_reprocess))),
                prospect_ids_to_reprocess)

                # Queue new lookup jobs
                lookup_jobs = await self._queue_lookup_jobs(prospect_ids_to_reprocess, organization_id)

                return {
                    "reprocessed": len(prospect_ids_to_reprocess),
                    "lookup_jobs_queued": len(lookup_jobs),
                    "prospects": [dict(p) for p in prospects]
                }

        except Exception as e:
            logger.error(f"Failed to reprocess prospects: {e}")
            return {"error": str(e)}

# Global service instance
ingestion_service = ProspectIngestionService()

# Backwards compatibility alias
IngestionService = ProspectIngestionService